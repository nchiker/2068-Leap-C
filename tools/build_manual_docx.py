#!/usr/bin/env python3
"""Build the Word edition of the Leap-C user's manual (docs/manual/*.md -> docs/2068_Leap-C_Users_Manual.docx).

Ported from the 2068-Leap V2 builder. The eight Markdown parts are concatenated in order; headings become
Word headings with bookmarks (so the links between parts work), tables, lists, code blocks and the pictures
are carried over. A ```c block (a program, test lines included -- they are ordinary C comments) is set
grey, a plain ``` block (screen output) green. Links like [x](05_concepts.md#memory) become Word
bookmarks named by the GitHub heading slug (`## #define` -> `define`; a repeated slug gets -1, -2 as on GitHub).

    python3 tools/build_manual_docx.py [output.docx]      (make manual)
    LEAPC_MANUAL_DIR=DIR   reads the parts (and ../images/ beside them) from DIR instead of docs/manual
"""
import os
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MANUAL = Path(os.environ.get("LEAPC_MANUAL_DIR") or ROOT / "docs" / "manual")
OUTPUT = ROOT / "docs" / "2068_Leap-C_Users_Manual.docx"
EDITION = "1.0.0-beta.2"
PARTS = ["01_introduction.md", "02_beginners_guide.md", "03_keywords.md", "04_library.md",
         "05_concepts.md", "06_libraries.md", "07_examples.md", "08_appendices.md"]
FLOW_PARTS = {"03_keywords.md", "04_library.md", "05_concepts.md"}   # A-Z entries flow on: no page break before each ##

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<![\w*])\*[^*\n]+\*(?![\w*])|\[[^]]+\]\([^)]+\))")
HEADING = re.compile(r"^(#{1,3})\s+(.+)$")
BOOKMARK_ID = 0


# --- helpers (word-level plumbing) -------------------------------------------------------------------

def markdown_anchor(text):
    """The GitHub-style anchor of a heading: `## #define` -> define."""
    text = re.sub(r"[`*_]", "", text).lower().strip()
    text = re.sub(r"[^\w\s-]", "", text)
    return re.sub(r"\s", "-", text)


def bookmark_name(anchor):
    """Map a Markdown anchor to a valid, stable Word bookmark name."""
    return "md_" + re.sub(r"[^A-Za-z0-9_]", "_", anchor)


def add_bookmark(paragraph, anchor):
    global BOOKMARK_ID
    BOOKMARK_ID += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(BOOKMARK_ID))
    start.set(qn("w:name"), bookmark_name(anchor))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(BOOKMARK_ID))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, label, target):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark_name(target.lstrip("#")))
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade(cell_or_paragraph, fill):
    props = (cell_or_paragraph._tc.get_or_add_tcPr() if hasattr(cell_or_paragraph, "_tc")
             else cell_or_paragraph._p.get_or_add_pPr())
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate.set(qn("w:dirty"), "true")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.extend((begin, text, separate, end))


def configure(document):
    """Page, styles (text, headings, code), header and footer."""
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 30, RGBColor(0x1F, 0x4E, 0x79)),
                              ("Heading 1", 20, RGBColor(0x1F, 0x4E, 0x79)),
                              ("Heading 2", 15, RGBColor(0x2E, 0x75, 0xB6)),
                              ("Heading 3", 12, RGBColor(0x3F, 0x3F, 0x3F))):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.right_indent = Inches(0.1)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)

    inline = document.styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
    inline.font.name = "Consolas"
    inline.font.size = Pt(9)
    inline.font.color.rgb = RGBColor(0x9C, 0x27, 0x2B)

    for sec in document.sections:
        header = sec.header.paragraphs[0]
        header.text = "2068-LEAP-C  •  TIMEX SINCLAIR 2068"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("2068-Leap-C User's Manual   •   ")
        add_field(footer, "PAGE")


# --- the manual --------------------------------------------------------------------------------------

def slugger():
    """A function giving each heading its anchor, GitHub style: a repeated slug gets -1, -2, ..."""
    seen = {}

    def unique(text):
        base = markdown_anchor(text)
        n = seen.get(base, 0)
        seen[base] = n + 1
        return base if n == 0 else "%s-%d" % (base, n)
    return unique


def heading_anchors():
    """The set of anchors the parts define (outside code blocks)."""
    unique, out = slugger(), set()
    for name in PARTS:
        in_code = False
        for line in (MANUAL / name).read_text(encoding="utf-8").splitlines():
            if line.startswith("```"):
                in_code = not in_code
            m = None if in_code else HEADING.match(line)
            if m:
                out.add(unique(m.group(2)))
    return out


def add_inline(paragraph, text, anchors):
    cursor = 0
    for match in INLINE.finditer(text):
        paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("`"):
            paragraph.add_run(token[1:-1]).style = "Code Inline"
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*"):
            first = len(paragraph.runs)
            add_inline(paragraph, token[1:-1], anchors)
            for run in paragraph.runs[first:]:
                run.italic = True
        else:
            label, target = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            anchor = target.partition("#")[2]
            if anchor and anchor in anchors:
                add_internal_hyperlink(paragraph, label, "#" + anchor)
            else:
                paragraph.add_run(label)
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def parse(document, lines, flow, unique, anchors):
    index, in_code, code_lines, info = 0, False, [], ""
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                shade(paragraph, "F2F4F7" if info == "c" else "E8F1E4")   # program grey, screen output green
                code_lines, in_code = [], False
            else:
                in_code, info = True, line[3:].strip()
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line or line == "---":
            index += 1
            continue
        image = re.match(r"^!\[([^]]*)\]\(([^)]+)\)\s*$", line)
        if image:
            path = MANUAL / image.group(2)
            if path.exists():
                document.add_picture(str(path), width=Inches(4.0))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(image.group(1))
                run.italic = True
                run.font.size = Pt(9)
            else:
                print("build_manual_docx: picture not found: %s" % path, file=sys.stderr)
            index += 1
            continue
        heading = HEADING.match(line)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_heading(level=level)
            if level == 1 or (level == 2 and not flow):
                paragraph.paragraph_format.page_break_before = True
            add_inline(paragraph, heading.group(2), anchors)
            add_bookmark(paragraph, unique(heading.group(2)))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([p.strip() for p in re.split(r"(?<!\\)\|", lines[index].strip().strip("|"))])
                index += 1
            rows.pop(1)
            table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
            table.style = "Light Shading Accent 1"
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    cell = table.cell(r, c)
                    add_inline(cell.paragraphs[0], value.replace("\\|", "|"), anchors)
                    if r == 0:
                        shade(cell, "1F4E79")
                        for run in cell.paragraphs[0].runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.bold = True
            document.add_paragraph()
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        number = re.match(r"^\d+\.\s+(.+)$", line)
        if bullet or number:
            item = (bullet or number).group(1)
            index += 1
            while index < len(lines) and re.match(r"^\s{2,}\S", lines[index]):
                item += " " + lines[index].strip()
                index += 1
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(paragraph, item, anchors)
            continue
        quote = line.startswith(">")
        parts = [line.lstrip("> ").strip() if quote else line.strip()]
        index += 1
        while index < len(lines):
            cand = lines[index]
            if (not cand or cand.startswith(("#", "```", "|", "- ", "* ", "![")) or cand == "---"
                    or re.match(r"^\d+\.\s", cand) or cand.startswith(">") != quote):
                break
            parts.append(cand.lstrip("> ").strip() if quote else cand.strip())
            index += 1
        paragraph = document.add_paragraph()
        if quote:
            paragraph.paragraph_format.left_indent = Inches(0.3)
        add_inline(paragraph, " ".join(parts), anchors)


def main():
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT
    missing = [n for n in PARTS if not (MANUAL / n).exists()]
    if missing:
        print("build_manual_docx: missing from %s: %s" % (MANUAL, ", ".join(missing)), file=sys.stderr)
        return 1
    anchors = heading_anchors()
    unique = slugger()
    document = Document()
    configure(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("2068-Leap-C -- User's Manual")
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("for the Timex Sinclair 2068")
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    document.add_paragraph("\n")
    ed = document.add_paragraph()
    ed.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ed.add_run("Edition %s • %s" % (EDITION, date.today().isoformat())).italic = True
    document.add_page_break()
    document.add_heading("Contents", level=1)
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u')
    note = document.add_paragraph("If page numbers are not visible, open the document in Word or LibreOffice "
                                  "and update the table of contents field.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)

    for name in PARTS:
        text = (MANUAL / name).read_text(encoding="utf-8").splitlines()
        parse(document, text, name in FLOW_PARTS, unique, anchors)
    document.core_properties.title = "2068-Leap-C User's Manual"
    document.core_properties.subject = "2068-Leap-C, edition %s" % EDITION
    document.core_properties.author = "2068-Leap-C Project"
    document.core_properties.keywords = "TS2068, Timex Sinclair, C, ROM, Leap-C"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print("Wrote %s" % output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
